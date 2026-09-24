"""
Tests for ai/text/document_text.py — PDF and .docx text extraction. No
mocking needed; these are pure file-format-parsing tests.
"""
import docx
import pytest
from reportlab.pdfgen import canvas

from ai.text.document_text import extract_text


@pytest.fixture
def sample_pdf(tmp_path):
    p = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(p))
    c.drawString(72, 720, "Epic 1: User Authentication")
    c.drawString(72, 700, "Users can log in with email and password.")
    c.showPage()
    c.drawString(72, 720, "Page two content.")
    c.save()
    return p


@pytest.fixture
def sample_docx(tmp_path):
    p = tmp_path / "sample.docx"
    d = docx.Document()
    d.add_paragraph("Change Request: Subsidiary Role")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Request Type"
    table.cell(0, 1).text = "Enhancement"
    table.cell(1, 0).text = "Classification"
    table.cell(1, 1).text = "Minor"
    d.add_paragraph("End of document.")
    d.save(str(p))
    return p


def test_extract_pdf_text_includes_page_markers_and_content(sample_pdf):
    text = extract_text(sample_pdf)
    assert "--- Page 1 ---" in text
    assert "--- Page 2 ---" in text
    assert "Epic 1: User Authentication" in text
    assert "Users can log in with email and password." in text
    assert "Page two content." in text
    # page 1's content appears before page 2's
    assert text.index("Epic 1") < text.index("Page two content.")


def test_extract_docx_text_preserves_paragraph_and_table_order(sample_docx):
    text = extract_text(sample_docx)
    assert "Change Request: Subsidiary Role" in text
    assert "Request Type: Enhancement" in text
    assert "Classification: Minor" in text
    assert "End of document." in text
    # document order preserved: intro paragraph, then table, then closing paragraph
    assert text.index("Change Request") < text.index("Request Type") < text.index("End of document.")


def test_extract_text_rejects_legacy_doc(tmp_path):
    p = tmp_path / "legacy.doc"
    p.write_bytes(b"not a real doc file")
    with pytest.raises(ValueError, match="Legacy .doc"):
        extract_text(p)


def test_extract_text_rejects_unsupported_suffix(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("plain text")
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(p)


def test_extract_text_raises_clear_value_error_for_corrupted_docx(tmp_path):
    """Confirmed directly against a real upload: a truncated .docx (an
    interrupted download — half the size of the original) raised python-
    docx's raw PackageNotFoundError, which the API surfaced as a misleading
    'transient AI response issue, please try again' 502 — actively wrong
    advice, since retrying the same corrupted file can never succeed. This
    must surface as a plain, actionable ValueError instead, which the API
    already maps to a proper 422."""
    p = tmp_path / "corrupted.docx"
    p.write_bytes(b"PK\x03\x04not a real docx, just the zip magic bytes and garbage")
    with pytest.raises(ValueError, match="corrupted|incomplete"):
        extract_text(p)


def test_extract_text_raises_clear_value_error_for_corrupted_pdf(tmp_path):
    p = tmp_path / "corrupted.pdf"
    p.write_bytes(b"not a real pdf at all")
    with pytest.raises(ValueError, match="corrupted|incomplete"):
        extract_text(p)
